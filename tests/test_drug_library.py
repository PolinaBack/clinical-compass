"""Technical fixtures are synthetic, never evidence about real medicines."""

import io
import json
import sqlite3
from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient
from app.main import create_app
from app.services import drug_library as lib
from app.storage import connect
from tests.helpers import login, token, TEST_HASH

INPUTS = Path(__file__).resolve().parents[1] / "docs" / "inputs" / "2026-10-21"


def upload(client, content, filename, **extra):
    return client.post(
        "/medicines/import/preview",
        data={"csrf": token(client, "/medicines"), "kind": "excel", **extra},
        files={"file": (filename, content)},
        follow_redirects=False,
    )


def confirm(client, iid):
    return client.post(
        f"/medicines/import/{iid}/confirm",
        data={"csrf": token(client, "/medicines"), "confirm": "yes"},
        follow_redirects=False,
    )


def iid(response):
    assert response.status_code == 303, response.text[:300]
    return int(response.headers["location"].split("/")[-1])


def workbook_bytes(w):
    f = io.BytesIO()
    w.save(f)
    return f.getvalue()


@pytest.mark.parametrize(
    "problem", ["empty_header", "duplicate_header", "extra_values"]
)
def test_excel_invalid_column_shape_reports_error(problem):
    w = openpyxl.load_workbook(INPUTS / "Библиотека препаратов.xlsx")
    s = w["Препараты"]
    if problem == "empty_header":
        s.cell(3, 7, "")
    elif problem == "duplicate_header":
        s.cell(3, 7, "МНН")
    else:
        s.cell(4, 19, "Значение без заголовка")
    _, _, errors, _ = lib.parse_excel(workbook_bytes(w))
    assert errors


def test_bootstrap_all_sheets_and_search(client):
    login(client, "doctor")
    with connect(client.app.state.db_path) as c:
        assert len(lib.current(c, "drug")) == 12
        assert len(lib.current(c, "pair")) == 11
        assert len(lib.current(c, "adverse")) == 468
        raw = lib.current(c, "raw")
        assert len(raw) == 14
        assert len(raw["library:Поля_инструкций"]["rows"][2]) == 22
        assert (
            sum(
                any(v is not None for v in r)
                for r in raw["library:Карточки_врача"]["rows"]
            )
            == 193
        )
        assert len(raw["adverse:Словарь_вопросов"]["rows"][3:]) == 163
    for query, name in [
        ("авиАНДР", "Авиандр"),
        ("маритупирдин", "Авиандр"),
        ("эсциталопрам", "Ципралекс"),
        ("серлифт", "Асентра"),
    ]:
        r = client.get("/api/medicines/suggest", params={"q": query})
        assert r.json()["results"][0]["name"] == name
        html = client.get("/medicines", params={"q": query}).text
        assert "Найдено: 1." in html and name in html
    assert "Найдено: 0." in client.get("/medicines?status=verified").text
    assert "Найдено: 4." in client.get("/medicines", params={"group": "СИОЗС"}).text
    for path in [
        "/medicines/aviandr",
        "/medicines/d12",
        "/reference/aviandr/archive",
        "/medicines/scenarios",
        "/medicines/import/template",
    ]:
        assert client.get(path).status_code == 200
    html = client.get("/medicines/aviandr").text
    assert "Подтверждённых данных в нашей базе нет" in html
    assert "Неподтверждённые исходные данные" in html


@pytest.mark.parametrize(
    "filename", ["Библиотека препаратов.xlsx", "Побочные эффекты по категориям.xlsx"]
)
def test_repeat_import_is_idempotent_and_explicit(client, filename):
    login(client, "doctor")
    with connect(client.app.state.db_path) as c:
        before = c.execute("SELECT count(*) FROM drug_entities").fetchone()[0]
    for _ in range(2):
        ident = iid(upload(client, (INPUTS / filename).read_bytes(), filename))
        with connect(client.app.state.db_path) as c:
            row = c.execute(
                "SELECT * FROM drug_imports WHERE id=?", (ident,)
            ).fetchone()
            report = json.loads(row["report"])
            assert not report["errors"] and report["duplicate_import"]
            assert all(v["status"] == "same" for v in report["changes"])
            assert row["committed_at"] is None
            assert (
                c.execute("SELECT count(*) FROM drug_entities").fetchone()[0] == before
            )
        assert confirm(client, ident).status_code == 303
        assert confirm(client, ident).status_code == 303
    with connect(client.app.state.db_path) as c:
        assert c.execute("SELECT count(*) FROM drug_entities").fetchone()[0] == before
        with pytest.raises(sqlite3.IntegrityError):
            c.execute("DELETE FROM drug_entities")


def test_update_history_conflict_and_no_delete(client):
    login(client, "doctor")
    w = openpyxl.load_workbook(INPUTS / "Библиотека препаратов.xlsx")
    w["Препараты"].cell(4, 11, "СИНТЕТИЧЕСКАЯ ПРАВКА ДЛЯ ТЕСТА")
    content = workbook_bytes(w)
    a = iid(upload(client, content, "update.xlsx"))
    b = iid(upload(client, content, "update.xlsx"))
    assert confirm(client, a).status_code == 303
    assert confirm(client, b).status_code == 409
    with connect(client.app.state.db_path) as c:
        assert len(lib.current(c, "drug")) == 12
        history = c.execute(
            "SELECT payload FROM drug_entities WHERE kind='drug' AND entity_key='D01' ORDER BY import_id"
        ).fetchall()
        assert len(history) == 2
        assert (
            json.loads(history[0][0])["raw"]["Предупреждения"]
            != "СИНТЕТИЧЕСКАЯ ПРАВКА ДЛЯ ТЕСТА"
        )
        assert (
            json.loads(history[1][0])["raw"]["Предупреждения"]
            == "СИНТЕТИЧЕСКАЯ ПРАВКА ДЛЯ ТЕСТА"
        )
        with pytest.raises(sqlite3.IntegrityError):
            c.execute("UPDATE drug_imports SET payload='{}' WHERE id=?", (a,))


@pytest.mark.parametrize(
    "problem", ["missing", "duplicate", "identity", "formula", "empty"]
)
def test_invalid_excel_blocked_with_report(client, problem):
    login(client, "doctor")
    w = openpyxl.load_workbook(INPUTS / "Библиотека препаратов.xlsx")
    if problem == "missing":
        w["Поля_инструкций"].cell(3, 18, "НЕВЕРНАЯ КОЛОНКА")
    elif problem == "duplicate":
        w["Препараты"].append(["D01", "Другое имя", "Другое вещество"])
    elif problem == "identity":
        w["Препараты"].cell(4, 3, "ДРУГОЙ МНН")
    elif problem == "formula":
        w["Препараты"].cell(4, 2, "=1+1")
    else:
        w["Препараты"].cell(4, 3, "")
    r = upload(client, workbook_bytes(w), "bad.xlsx")
    if problem == "formula":
        assert r.status_code == 422
    else:
        ident = iid(r)
        with connect(client.app.state.db_path) as c:
            assert json.loads(
                c.execute(
                    "SELECT report FROM drug_imports WHERE id=?", (ident,)
                ).fetchone()[0]
            )["errors"]
        assert confirm(client, ident).status_code == 422
    with connect(client.app.state.db_path) as c:
        assert len(lib.current(c, "drug")) == 12


def test_combinations_two_three_and_unknown(client):
    login(client, "doctor")
    html = client.get(
        "/medicines/combinations", params={"drug": ["D01", "D02", "D04"]}
    ).text
    assert html.count('class="card medicine-pair"') == 3
    assert html.count("Подтверждённое взаимодействие не найдено в нашей базе.") == 3
    assert html.count("Есть неподтверждённая запись Excel") == 2
    assert "Это не означает, что взаимодействий нет." in html
    assert (
        "Проверено пар: 1."
        in client.get(
            "/medicines/combinations", params={"drug": ["D01", "D02", "D01"]}
        ).text
    )
    assert client.get("/medicines/combinations?drug=D01").status_code == 422
    assert client.get("/medicines/combinations?drug=D01&drug=D99").status_code == 422


def official_fixture(client):
    """Fake content marked as synthetic; never used in real DB or shipped as evidence."""
    w = openpyxl.Workbook()
    s = w.active
    s.title = "Сведения"
    s.append(["Поле", "Текст", "Раздел/страница"])
    s.append(["warnings", "СИНТЕТИЧЕСКОЕ ПРЕДУПРЕЖДЕНИЕ — тест", "Тестовый раздел 4"])
    s = w.create_sheet("Взаимодействия")
    s.append(["Препарат A", "Препарат B", "Уровень", "Объяснение", "Раздел/страница"])
    s.append(
        [
            "D01",
            "D02",
            "caution",
            "СИНТЕТИЧЕСКОЕ ВЗАИМОДЕЙСТВИЕ — тест",
            "Тестовый раздел 5",
        ]
    )
    meta = dict(
        kind="official",
        drug_code="D01",
        source_url="https://grls.rosminzdrav.ru/TEST-NOT-A-REAL-SOURCE",
        version="СИНТЕТИЧЕСКАЯ ТЕСТОВАЯ ВЕРСИЯ",
        registration="TEST-RU",
        product_form="СИНТЕТИЧЕСКАЯ ФОРМА",
        checked_on=date.today().isoformat(),
        valid_until=(date.today() + timedelta(days=30)).isoformat(),
    )
    return workbook_bytes(w), meta


def review(client, ident, action="verify"):
    return client.post(
        f"/medicines/import/{ident}/review",
        data={
            "csrf": token(client, "/medicines"),
            "action": action,
            "verified": "yes",
            "note": "ТЕСТОВАЯ СВЕРКА — не медицинские сведения",
        },
        follow_redirects=False,
    )


def test_official_review_provenance_withdrawal_expiry(client):
    login(client, "doctor")
    content, meta = official_fixture(client)
    ident = iid(upload(client, content, "synthetic.xlsx", **meta))
    assert confirm(client, ident).status_code == 303
    assert "СИНТЕТИЧЕСКОЕ ПРЕДУПРЕЖДЕНИЕ" not in client.get("/medicines/aviandr").text
    assert review(client, ident).status_code == 303
    assert "СИНТЕТИЧЕСКОЕ ПРЕДУПРЕЖДЕНИЕ" in client.get("/medicines/aviandr").text
    html = client.get("/medicines/combinations?drug=D01&drug=D02&drug=D04").text
    assert html.count("СИНТЕТИЧЕСКОЕ ВЗАИМОДЕЙСТВИЕ") == 1
    assert (
        "Тестовый раздел 5" in html
        and html.count("Подтверждённое взаимодействие не найдено") == 2
    )
    assert iid(upload(client, content, "synthetic.xlsx", **meta)) == ident
    assert review(client, ident, "withdraw").status_code == 303
    assert "СИНТЕТИЧЕСКОЕ ПРЕДУПРЕЖДЕНИЕ" not in client.get("/medicines/aviandr").text
    assert "отозвал проверку" in client.get(f"/medicines/import/{ident}").text
    meta["valid_until"] = (date.today() - timedelta(days=1)).isoformat()
    assert upload(client, content, "synthetic.xlsx", **meta).status_code == 422


def test_new_source_retires_old_without_reviving_on_withdrawal(client):
    login(client, "doctor")
    content, meta = official_fixture(client)
    first = iid(upload(client, content, "synthetic.xlsx", **meta))
    confirm(client, first)
    review(client, first)
    meta["version"] = "СИНТЕТИЧЕСКАЯ НОВАЯ ВЕРСИЯ"
    second = iid(upload(client, content, "synthetic.xlsx", **meta))
    confirm(client, second)
    with connect(client.app.state.db_path) as c:
        assert [s["id"] for s in lib.sources(c)] == [first]
    review(client, second)
    with connect(client.app.state.db_path) as c:
        assert [s["id"] for s in lib.sources(c)] == [second]
    review(client, second, "withdraw")
    with connect(client.app.state.db_path) as c:
        assert not lib.sources(c)
        assert len(lib.sources(c, include_inactive=True)) == 2


def test_verification_expires_and_cannot_publish_spreadsheet(client, monkeypatch):
    login(client, "doctor")
    content, meta = official_fixture(client)
    ident = iid(upload(client, content, "synthetic.xlsx", **meta))
    confirm(client, ident)
    denied = client.post(
        f"/medicines/import/{ident}/review",
        data={
            "csrf": token(client, "/medicines"),
            "action": "verify",
            "note": "Синтетическая запись без подтверждения",
        },
    )
    assert denied.status_code == 422
    assert review(client, 1).status_code == 422
    review(client, ident)

    class FutureDate(date):
        @classmethod
        def today(cls):
            return date.fromisoformat(meta["valid_until"]) + timedelta(days=1)

    monkeypatch.setattr(lib, "date", FutureDate)
    with connect(client.app.state.db_path) as c:
        assert not lib.sources(c)
    assert "СИНТЕТИЧЕСКОЕ ПРЕДУПРЕЖДЕНИЕ" not in client.get("/medicines/aviandr").text


@pytest.mark.parametrize(
    "url",
    [
        "http://grls.rosminzdrav.ru/doc",
        "https://grls.rosminzdrav.ru.evil.test/doc",
        "https://www.vidal.ru/doc",
        "https://grls.rosminzdrav.ru/",
        "javascript:alert(1)",
    ],
)
def test_official_source_validation(client, url):
    login(client, "doctor")
    content, meta = official_fixture(client)
    meta["source_url"] = url
    assert upload(client, content, "synthetic.xlsx", **meta).status_code == 422


def test_access_csrf_and_import_owner(client):
    login(client, "doctor")
    ident = iid(
        upload(
            client, (INPUTS / "Библиотека препаратов.xlsx").read_bytes(), "library.xlsx"
        )
    )
    assert (
        client.post(
            f"/medicines/import/{ident}/confirm", data={"confirm": "yes"}
        ).status_code
        == 403
    )
    login(client, "patient1")
    for path in [
        "/medicines",
        "/medicines/aviandr",
        "/medicines/combinations",
        "/medicines/import",
        "/medicines/import/template",
        f"/medicines/import/{ident}",
        f"/medicines/import/{ident}/source",
        "/api/medicines/suggest?q=ави",
    ]:
        assert client.get(path).status_code == 403
    for path in [
        "/medicines/import/preview",
        f"/medicines/import/{ident}/confirm",
        f"/medicines/import/{ident}/review",
    ]:
        assert client.post(path, data={"csrf": token(client)}).status_code == 403
    assert client.get("/help-now").status_code == 200
    login(client, "doctor")
    with connect(client.app.state.db_path) as c:
        c.execute("INSERT INTO doctors(id,name) VALUES(2,'Другой врач')")
        c.execute(
            "INSERT INTO accounts(username,password_hash,role,doctor_id) VALUES('doctor-other',?,'doctor',2)",
            (TEST_HASH,),
        )
    client.post("/logout", data={"csrf": token(client)})
    from tests.helpers import TEST_PASSWORD

    client.post(
        "/login",
        data={
            "username": "doctor-other",
            "password": TEST_PASSWORD,
            "portal": "specialist",
            "csrf": token(client, "/login/specialist"),
        },
    )
    assert confirm(client, ident).status_code == 403


def test_pdf_and_reboot_no_data_loss(tmp_path):
    path = tmp_path / "old.db"
    with sqlite3.connect(path) as c:
        c.executescript(
            """CREATE TABLE doctors(id INTEGER PRIMARY KEY,name TEXT NOT NULL);
        INSERT INTO doctors VALUES(1,'Старый врач');
        CREATE TABLE patients(id INTEGER PRIMARY KEY,name TEXT NOT NULL,doctor_id INTEGER NOT NULL DEFAULT 1 REFERENCES doctors(id));
        INSERT INTO patients VALUES(1,'Старый пациент',1);
        CREATE TABLE diary(id INTEGER PRIMARY KEY,patient_id INTEGER,day TEXT NOT NULL,wellbeing INTEGER NOT NULL,sleep REAL NOT NULL,adherence TEXT NOT NULL,note TEXT NOT NULL,contact INTEGER NOT NULL);
        INSERT INTO diary VALUES(42,1,'2020-01-01',0,8.5,'Нет','Старая запись',1);"""
        )
        old = c.execute("SELECT * FROM diary").fetchall()
    for _ in range(2):
        with TestClient(create_app(path)) as client:
            login(client, "doctor")
            _, meta = official_fixture(client)
            meta.update(
                warnings="СИНТЕТИЧЕСКАЯ ВЫДЕРЖКА", warnings_locator="Тестовая стр. 1"
            )
            r = upload(
                client, b"%PDF-1.4\n% SYNTHETIC TEST ONLY", "synthetic.pdf", **meta
            )
            ident = iid(r)
            assert confirm(client, ident).status_code == 303
            assert client.get(f"/medicines/import/{ident}/source").content.startswith(
                b"%PDF-"
            )
    with connect(path) as c:
        assert [tuple(r)[:8] for r in c.execute("SELECT * FROM diary")] == old
        assert c.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        assert not c.execute("PRAGMA foreign_key_check").fetchall()
        assert len(lib.current(c, "drug")) == 12
