"""Physician-only reference and comparison routes; no treatment mutations."""

import io
import json
from urllib.parse import urlencode, urlsplit, parse_qs
from fastapi import Request, HTTPException
from fastapi.responses import RedirectResponse, Response
import openpyxl
from app.storage import audit, now
from app.services import drug_library as lib
from app.services import medicine_reference as editorial
from app.services.navigation import local_destination


def register_drug_library(app, db, account, doctor_access, csrf_check, render):
    def medicine_return(request, value, fallback="/medicines"):
        destination = local_destination(request, value)
        return (
            destination
            if destination
            and (
                destination.split("?", 1)[0].split("#", 1)[0] == "/medicines"
                or destination.startswith("/medicines/")
            )
            else fallback
        )

    def get_import(c, iid, user, owner=True):
        row = c.execute("SELECT * FROM drug_imports WHERE id=?", (iid,)).fetchone()
        if not row:
            raise HTTPException(404, "Импорт не найден")
        if owner and row["actor_id"] not in (None, user["id"]):
            raise HTTPException(403, "Это предпросмотр другого специалиста")
        return row

    @app.get("/medicines")
    def library(request: Request, q: str = "", group: str = "", status: str = ""):
        doctor_access(request)
        with db() as c:
            drugs = lib.search(c, q[:200], group, status)
            groups = sorted({d["group"] for d in lib.current(c, "drug").values()})
        return render(
            request,
            "medicines.html",
            medicines=drugs,
            q=q[:200],
            groups=groups,
            group=group,
            status=status,
            comparison=request.session.get("medicine_comparison", []),
            catalogue_return=request.url.path
            + ("?" + request.url.query if request.url.query else ""),
        )

    @app.get("/api/medicines/suggest")
    def suggest(request: Request, q: str = ""):
        doctor_access(request)
        with db() as c:
            return {
                "results": [
                    {
                        "code": d["code"],
                        "name": d["name"],
                        "inn": d["inn"],
                        "url": "/medicines/" + d["slug"],
                    }
                    for d in lib.search(c, q[:200])[:12]
                ]
            }

    @app.get("/medicines/import")
    def import_page(request: Request):
        doctor_access(request)
        with db() as c:
            imports = c.execute(
                "SELECT id,kind,filename,created_at,committed_at,actor_id FROM drug_imports ORDER BY id DESC LIMIT 100"
            ).fetchall()
            drugs = lib.search(c)
        return render(
            request,
            "medicine_import.html",
            imports=imports,
            drugs=drugs,
            fields=lib.FIELDS,
        )

    @app.post("/medicines/import/preview")
    async def import_preview(request: Request):
        doctor_access(request)
        user = account(request)
        form = await request.form(max_files=1, max_fields=60)
        csrf_check(request, str(form.get("csrf", "")))
        upload = form.get("file")
        if not upload or not hasattr(upload, "read"):
            lib.fail("Выберите файл.")
        content = await upload.read(lib.MAX_FILE + 1)
        await upload.close()
        if not content or len(content) > lib.MAX_FILE:
            lib.fail("Файл пуст или превышает 8 МБ.")
        with db() as c:
            official = (
                lib.official_payload(c, content, upload.filename or "", form)
                if form.get("kind") == "official"
                else None
            )
            if official is None and not (upload.filename or "").lower().endswith(
                ".xlsx"
            ):
                lib.fail("Для библиотек нужен XLSX.")
            iid = lib.preview(
                c, upload.filename or "file", content, user["id"], official
            )
            audit(c, user["id"], "drug_import_preview", None, iid)
        return RedirectResponse(f"/medicines/import/{iid}", status_code=303)

    @app.get("/medicines/import/template")
    def template(request: Request):
        doctor_access(request)
        w = openpyxl.Workbook()
        s = w.active
        s.title = "Сведения"
        s.append(["Поле", "Текст", "Раздел/страница"])
        p = w.create_sheet("Взаимодействия")
        p.append(
            ["Препарат A", "Препарат B", "Уровень", "Объяснение", "Раздел/страница"]
        )
        guide = w.create_sheet("Как заполнить")
        for r in [
            [
                "Назначение",
                "Ручная выписка из официальной инструкции, а не автоматическая выгрузка ГРЛС",
            ],
            ["Поля", ", ".join(lib.FIELDS)],
            ["Уровни", ", ".join(lib.SEVERITIES)],
            [
                "Препараты",
                "Коды D01…D12; одна сторона пары должна совпадать с выбранным препаратом документа",
            ],
            [
                "Неизвестное",
                "Не заполняйте поле, если в источнике нет подтверждённого текста.",
            ],
            [
                "Источники",
                "В форме импорта обязательны точная ссылка ГРЛС, редакция, РУ, форма и даты.",
            ],
        ]:
            guide.append(r)
        for sheet in w:
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = 35
        data = io.BytesIO()
        w.save(data)
        return Response(
            data.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": 'attachment; filename="official-reference-template.xlsx"'
            },
        )

    @app.get("/medicines/import/{iid}")
    def show_import(request: Request, iid: int):
        doctor_access(request)
        with db() as c:
            row = get_import(c, iid, account(request), owner=False)
            report = json.loads(row["report"])
            payload = json.loads(row["payload"])
            events = c.execute(
                """SELECT e.*,a.username FROM drug_source_events e JOIN accounts a ON a.id=e.actor_id
                               WHERE import_id=? ORDER BY e.id DESC""",
                (iid,),
            ).fetchall()
        return render(
            request,
            "medicine_preview.html",
            item=row,
            report=report,
            payload=payload,
            events=events,
            fields=lib.FIELDS,
            severities=lib.SEVERITIES,
            own=row["actor_id"] == account(request)["id"],
        )

    @app.post("/medicines/import/{iid}/confirm")
    async def confirm_import(request: Request, iid: int):
        doctor_access(request)
        user = account(request)
        form = await request.form()
        csrf_check(request, str(form.get("csrf", "")))
        if form.get("confirm") != "yes":
            lib.fail("Подтвердите добавление/обновление без удаления данных.")
        with db() as c:
            c.execute("BEGIN IMMEDIATE")
            get_import(c, iid, user)
            lib.commit(c, iid)
            audit(c, user["id"], "drug_import_committed", None, iid)
        return RedirectResponse(f"/medicines/import/{iid}", status_code=303)

    @app.get("/medicines/import/{iid}/source")
    def source(request: Request, iid: int):
        doctor_access(request)
        with db() as c:
            row = get_import(c, iid, account(request), owner=False)
        mime = (
            "application/pdf"
            if row["filename"].lower().endswith(".pdf")
            else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        return Response(
            row["content"],
            media_type=mime,
            headers={
                "Content-Disposition": f'attachment; filename="source-{iid}{".pdf" if mime=="application/pdf" else ".xlsx"}"',
                "X-Content-Type-Options": "nosniff",
                "Cache-Control": "no-store",
            },
        )

    @app.post("/medicines/import/{iid}/review")
    async def review(request: Request, iid: int):
        doctor_access(request)
        user = account(request)
        form = await request.form()
        csrf_check(request, str(form.get("csrf", "")))
        action = str(form.get("action", ""))
        note = str(form.get("note", "")).strip()
        if action not in ("verify", "withdraw") or not 10 <= len(note) <= 3000:
            lib.fail("Укажите действие и обоснование от 10 до 3000 символов.")
        with db() as c:
            c.execute("BEGIN IMMEDIATE")
            row = get_import(c, iid, user, owner=False)
            if row["kind"] != "official" or not row["committed_at"]:
                lib.fail("Можно проверять только импортированный официальный документ.")
            meta = json.loads(row["payload"])["meta"]
            if action == "verify":
                if form.get("verified") != "yes":
                    lib.fail(
                        "Подтвердите сверку точного продукта, РУ, формы и действующей версии в ГРЛС."
                    )
                if meta["valid_until"] < lib.date.today().isoformat():
                    lib.fail("Срок сверки истёк; загрузите новую версию.")
            c.execute(
                "INSERT INTO drug_source_events(import_id,action,actor_id,note,created_at) VALUES(?,?,?,?,?)",
                (iid, action, user["id"], note, now()),
            )
            audit(c, user["id"], "drug_source_" + action, None, iid)
        return RedirectResponse(f"/medicines/import/{iid}", status_code=303)

    @app.post("/medicines/combinations/selection")
    async def select_comparison(request: Request):
        doctor_access(request)
        form = await request.form(max_files=0, max_fields=40)
        csrf_check(request, str(form.get("csrf", "")))
        selected = list(dict.fromkeys(form.getlist("selected")))
        action, code = str(form.get("action", "")), str(form.get("code", ""))
        with db() as c:
            drugs = lib.current(c, "drug")
            if len(selected) > 20 or any(d not in drugs for d in selected):
                lib.fail("Некорректный набор для проверки")
            if action == "add" and code in drugs:
                selected = list(dict.fromkeys(selected + [code]))
            elif action == "remove" and code in drugs:
                selected = [d for d in selected if d != code]
            elif action == "clear":
                selected = []
            else:
                lib.fail("Неизвестное действие с набором")
            if len(selected) > 20:
                lib.fail("Допустимо до 20 препаратов")
            pid = str(form.get("draft_patient", ""))
            if pid:
                from app.services.medicine_drafts import patient

                try:
                    patient(c, account(request), int(pid))
                except ValueError:
                    lib.fail("Некорректная учебная карта")
        if not pid:
            request.session["medicine_comparison"] = selected
        target = medicine_return(
            request, str(form.get("return_to", "")), "/medicines/combinations"
        )
        if target.split("?", 1)[0] == "/medicines/combinations":
            params = [("drug", d) for d in selected] or [("empty", "1")]
            if pid:
                params.append(("draft_patient", pid))
            if form.get("picker") == "1":
                params.append(("picker", "1"))
            target = (
                "/medicines/combinations?"
                + urlencode(params)
                + ("#medicine-picker" if form.get("picker") == "1" else "#comparison")
            )
        return RedirectResponse(target, status_code=303)

    @app.get("/medicines/combinations")
    def combinations_page(request: Request, draft_patient: int | None = None):
        doctor_access(request)
        codes = list(
            dict.fromkeys(
                request.query_params.getlist("drug")
                if "drug" in request.query_params
                else (
                    []
                    if request.query_params.get("empty")
                    else request.session.get("medicine_comparison", [])
                )
            )
        )
        results = None
        with db() as c:
            drugs = lib.search(c)
            if request.query_params.get("all") == "1":
                codes = [d["code"] for d in drugs]
            if draft_patient is not None:
                from app.services.medicine_drafts import patient

                patient(c, account(request), draft_patient)
            if len(codes) > 20 or any(
                code not in {d["code"] for d in drugs} for code in codes
            ):
                lib.fail("Неизвестный препарат в наборе")
            if len(codes) >= 2:
                results = lib.check_pairs(c, codes)
        target = "/medicines/combinations?" + urlencode(
            ([("drug", code) for code in codes] or [("empty", "1")])
            + ([("draft_patient", draft_patient)] if draft_patient is not None else [])
        )
        for d in drugs:
            d["card_url"] = (
                "/medicines/" + d["slug"] + "?" + urlencode(dict(return_to=target))
            )
        if results:
            links = {d["code"]: d["card_url"] for d in drugs}
            for result in results:
                for side in ("a", "b"):
                    result[side]["card_url"] = links[result[side]["code"]]
        return render(
            request,
            "medicine_combinations.html",
            drugs=drugs,
            selected=codes,
            results=results,
            severities=lib.SEVERITIES,
            draft_patient=draft_patient,
            picker_open=request.query_params.get("picker") == "1" or len(codes) < 2,
            comparison_return=target,
        )

    @app.get("/medicines/scenarios")
    def scenarios(request: Request):
        doctor_access(request)
        return render(request, "medicine_scenarios.html")

    @app.get("/medicines/{slug}")
    def drug_card(request: Request, slug: str):
        doctor_access(request)
        code = "D01" if slug == "aviandr" else slug.upper()
        with db() as c:
            drugs = lib.current(c, "drug")
            drug = drugs.get(code)
            if not drug:
                raise HTTPException(404, "Лекарство не найдено")
            drug = editorial.enrich(drug)
            refs = lib.sources(c, code, include_inactive=True)
            facts = {field: [] for field in lib.FIELDS}
            for s in refs:
                if s["active"]:
                    for f in s["meta"]["facts"]:
                        facts[f["field"]].append(dict(**f, source=s))
            # Unverified archives are available explicitly, never labelled as clinical facts.
            raw = []
            for key, value in lib.current(c, "raw").items():
                if key.startswith("library:"):
                    rows = value["rows"]
                    matching = [
                        (i + 1, r)
                        for i, r in enumerate(rows)
                        if r
                        and (str(r[0]) == code or str(r[0]).startswith(code + " •"))
                    ]
                    if matching:
                        raw.append(
                            dict(
                                sheet=key.split(":", 1)[1],
                                rows=matching,
                                import_id=value["_import"],
                            )
                        )
            adverse = [
                v
                for v in lib.current(c, "adverse").values()
                if str(v["raw"].get("ID препарата")) == code
            ]
            history = c.execute(
                "SELECT e.import_id,i.filename,i.created_at,e.payload FROM drug_entities e JOIN drug_imports i ON i.id=e.import_id WHERE e.kind='drug' AND e.entity_key=? ORDER BY e.import_id DESC",
                (code,),
            ).fetchall()
            return_target = medicine_return(
                request, request.query_params.get("return_to")
            )
            comparison = request.session.get("medicine_comparison", [])
            comparison_link = "/medicines/combinations"
            comparison_patient = None
            if urlsplit(return_target).path == "/medicines/combinations":
                params = parse_qs(urlsplit(return_target).query)
                comparison = params.get("drug", [])
                if len(comparison) > 20 or any(
                    item not in drugs for item in comparison
                ):
                    lib.fail("Некорректный набор для возврата к проверке")
                comparison_link = return_target
                if params.get("draft_patient"):
                    from app.services.medicine_drafts import patient

                    try:
                        comparison_patient = int(params["draft_patient"][0])
                    except ValueError:
                        lib.fail("Некорректная учебная карта")
                    patient(c, account(request), comparison_patient)
        return render(
            request,
            "medicine_card.html",
            drug=drug,
            refs=refs,
            facts=facts,
            fields=lib.FIELDS,
            raw=raw,
            adverse=adverse,
            history=history,
            comparison=comparison,
            comparison_link=comparison_link,
            comparison_patient=comparison_patient,
            card_return=return_target,
            self_return=(
                comparison_link
                if comparison_link != "/medicines/combinations"
                else request.url.path
                + ("?" + request.url.query if request.url.query else "")
            ),
        )
