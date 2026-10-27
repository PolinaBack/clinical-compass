"""Versioned reference data. Never reads or changes a patient's treatment."""

import hashlib
import io
import json
import re
import unicodedata
import zipfile
from datetime import date, timedelta
from itertools import combinations
from pathlib import Path
from urllib.parse import urlsplit

import openpyxl
from fastapi import HTTPException
from app.storage import now

MAX_FILE = 8 * 1024 * 1024
FIELDS = {
    "indications": "Показания",
    "contraindications": "Противопоказания",
    "warnings": "Предупреждения",
    "adverse": "Нежелательные реакции",
    "interactions": "Взаимодействия",
    "doses": "Способ применения и дозы",
    "manufacturer": "Производитель и держатель РУ",
    "storage": "Хранение и срок годности",
    "pharmacovigilance": "Фармаконадзор и контакты",
}
OFFICIAL_HOSTS = {"grls.rosminzdrav.ru", "grls.minzdrav.gov.ru"}
SEVERITIES = {
    "contraindicated": "Противопоказано по источнику",
    "high": "Высокий уровень внимания",
    "caution": "Требует осторожности",
    "information": "Информация из источника",
}


def fail(message):
    raise HTTPException(422, message)


def dump(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def normalize(value):
    return (
        unicodedata.normalize("NFKC", str(value or ""))
        .casefold()
        .replace("ё", "е")
        .strip()
    )


def safe_url(value, official=False):
    try:
        p = urlsplit(str(value))
        good = (
            p.scheme == "https"
            and p.hostname
            and not p.username
            and not p.password
            and p.port in (None, 443)
        )
        if official:
            good = good and p.hostname in OFFICIAL_HOSTS and p.path not in ("", "/")
        return bool(good)
    except ValueError:
        return False


def spreadsheet(content):
    if not content or len(content) > MAX_FILE:
        fail("Файл пуст или превышает 8 МБ.")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            infos = z.infolist()
            if len(infos) > 2000 or sum(i.file_size for i in infos) > 32 * 1024 * 1024:
                fail("Слишком большой распакованный Excel.")
            if any(
                "vbaProject" in i.filename or i.filename.startswith("xl/externalLinks/")
                for i in infos
            ):
                fail("Макросы и внешние связи не поддерживаются.")
        w = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        sheets = {}
        count = 0
        for s in w:
            if (s.max_row or 0) > 10000 or (s.max_column or 0) > 50:
                fail("Лист превышает 10000 строк или 50 колонок.")
            rows = []
            for row_number, row in enumerate(s, 1):
                if row_number > 10000 or len(row) > 50:
                    fail("Лист превышает 10000 строк или 50 колонок.")
                count += len(row)
                if count > 200000:
                    fail("Слишком много ячеек.")
                values = []
                for cell in row:
                    if cell.data_type == "f":
                        fail(
                            f"Формула в {s.title}!{cell.coordinate}: загрузите значения вместо формул."
                        )
                    v = cell.value
                    if isinstance(v, (date,)):
                        v = v.isoformat()
                    if v is not None and not isinstance(v, (str, int, float, bool)):
                        v = str(v)
                    if len(str(v or "")) > 20000:
                        fail("Текст ячейки превышает 20000 символов.")
                    values.append(v)
                rows.append(values)
            sheets[s.title] = rows
        w.close()
        return sheets
    except HTTPException:
        raise
    except Exception:
        fail("Не удалось прочитать XLSX. Проверьте формат и отсутствие пароля.")


def table(sheets, name, header, required, errors):
    rows = sheets.get(name)
    if not rows:
        errors.append(f"Отсутствует лист «{name}».")
        return []
    pos = next((i for i, r in enumerate(rows[:5]) if r and r[0] == header), None)
    if pos is None or any(x not in rows[pos] for x in required):
        errors.append(
            f'Лист «{name}»: неверные обязательные колонки: {", ".join(required)}.'
        )
        return []
    headings = rows[pos]
    while headings and headings[-1] is None:
        headings = headings[:-1]
    if any(not isinstance(h, str) or not h.strip() for h in headings) or len(
        set(headings)
    ) != len(headings):
        errors.append(f"Лист «{name}»: пустые или повторяющиеся заголовки колонок.")
        return []
    if any(any(v is not None for v in r[len(headings) :]) for r in rows[pos + 1 :]):
        errors.append(f"Лист «{name}»: есть значения в колонках без заголовков.")
        return []
    return [
        (i + 1, {h: row[j] if j < len(row) else None for j, h in enumerate(headings)})
        for i, row in enumerate(rows)
        if i > pos and any(v is not None for v in row)
    ]


def current(c, kind):
    rows = c.execute(
        """SELECT e.* FROM drug_entities e JOIN
       (SELECT entity_key,MAX(import_id) latest FROM drug_entities WHERE kind=? GROUP BY entity_key) x
       ON e.entity_key=x.entity_key AND e.import_id=x.latest WHERE e.kind=?""",
        (kind, kind),
    ).fetchall()
    return {
        r["entity_key"]: {**json.loads(r["payload"]), "_import": r["import_id"]}
        for r in rows
    }


def parse_excel(content):
    sheets = spreadsheet(content)
    errors, warnings, entities = [], [], []
    if "Препараты" in sheets:
        kind = "library"
        expected = [
            "ПЕРЕД НАЧАЛОМ",
            "Препараты",
            "Авиандр_комбинации",
            "Сценарий_демо",
            "Поля_инструкций",
            "Карточки_врача",
            "Источники_и_статусы",
            "НЯ_468_записей",
        ]
        drugs = table(
            sheets,
            "Препараты",
            "ID",
            [
                "ID",
                "Препарат (пример)",
                "МНН",
                "Группа",
                "ATX",
                "Форма / дозировка",
                "Другие бренды",
            ],
            errors,
        )
        seen = {}
        for row, raw in drugs:
            code = str(raw.get("ID") or "").strip()
            if (
                not re.fullmatch(r"D\d{2,5}", code)
                or not raw.get("МНН")
                or not raw.get("Препарат (пример)")
            ):
                errors.append(
                    f"Препараты, строка {row}: обязательны ID вида D01, название и МНН."
                )
                continue
            value = dict(
                code=code,
                name=raw["Препарат (пример)"],
                inn=raw["МНН"],
                substance=raw["МНН"],
                group=raw["Группа"] or "",
                atc=raw["ATX"] or "",
                form=raw["Форма / дозировка"] or "",
                aliases=raw["Другие бренды"] or "",
                raw=raw,
                row=row,
            )
            if code in seen:
                if seen[code] != raw:
                    errors.append(
                        f"Препараты, строка {row}: конфликтующий дубликат {code}."
                    )
                else:
                    warnings.append(f"Повтор {code} в строке {row}, запись будет одна.")
                continue
            seen[code] = raw
            entities.append(("drug", code, value))
        pairs = table(
            sheets,
            "Авиандр_комбинации",
            "Сравниваем с Авиандром",
            [
                "Сравниваем с Авиандром",
                "МНН",
                "Уровень внимания",
                "Почему",
                "Источник B",
                "Источник A",
            ],
            errors,
        )
        for row, raw in pairs:
            matches = [
                k
                for k, r in seen.items()
                if normalize(r["МНН"]) == normalize(raw["МНН"])
            ]
            if len(matches) != 1 or matches[0] == "D01" or "D01" not in seen:
                errors.append(
                    f"Сочетания, строка {row}: не удалось однозначно связать препараты."
                )
                continue
            key = "|".join(sorted(["D01", matches[0]]))
            entities.append(
                (
                    "pair",
                    key,
                    dict(
                        raw=raw,
                        row=row,
                        level=raw["Уровень внимания"],
                        explanation=raw["Почему"],
                        urls=[raw["Источник A"], raw["Источник B"]],
                    ),
                )
            )
        table(
            sheets,
            "Поля_инструкций",
            "ID",
            ["ID", "Бренд", "МНН", "Показания", "Противопоказания", "Дата проверки"],
            errors,
        )
        table(
            sheets,
            "Карточки_врача",
            "Препарат (МНН)",
            [
                "Препарат (МНН)",
                "Раздел карточки",
                "Содержание / сведения о конкретной форме",
                "Источник и статус",
            ],
            errors,
        )
    elif "Реакции_по_препаратам" in sheets:
        kind = "adverse"
        expected = [
            "ОПИСАНИЕ",
            "Реакции_по_препаратам",
            "Категории",
            "Словарь_вопросов",
            "Сигналы_безопасности",
            "Для_разработчика",
        ]
        reactions = table(
            sheets,
            "Реакции_по_препаратам",
            "ID препарата",
            [
                "ID препарата",
                "МНН",
                "Категория",
                "Медицинский термин",
                "Частота по источнику",
            ],
            errors,
        )
        for row, raw in reactions:
            if any(
                not raw.get(k)
                for k in ("ID препарата", "МНН", "Категория", "Медицинский термин")
            ):
                errors.append(
                    f"Реакции, строка {row}: обязательны ID, МНН, категория и термин."
                )
                continue
            entities.append(
                (
                    "adverse",
                    f'{raw["ID препарата"]}|{normalize(raw["Категория"])}|{normalize(raw["Медицинский термин"])}',
                    dict(raw=raw, row=row),
                )
            )
        table(
            sheets,
            "Словарь_вопросов",
            "ID симптома",
            ["ID симптома", "Категория", "Клиническое обозначение"],
            errors,
        )
    else:
        fail(
            "Неизвестная книга: нужны исходная библиотека препаратов или книга побочных эффектов."
        )
    schema = json.loads(
        (
            Path(__file__).resolve().parents[1]
            / "references"
            / "drug_import_schema.json"
        ).read_text(encoding="utf-8")
    )[kind]
    for sheet, columns in schema.items():
        if sheet in sheets:
            table(sheets, sheet, columns[0], columns, errors)
    for name in expected:
        if name not in sheets and not any(f"«{name}»" in e for e in errors):
            errors.append(f"Отсутствует лист «{name}».")
    # Every cell is retained, including warnings, unknowns, dictionary and editorial cards.
    entities.extend(
        ("raw", f"{kind}:{name}", dict(rows=rows)) for name, rows in sheets.items()
    )
    if not errors and not any(e[0] in ("drug", "adverse") for e in entities):
        errors.append("Нет записей для импорта.")
    return kind, dict(sheets=sheets, entities=entities), errors, warnings


def preview(c, filename, content, actor_id, official=None):
    if official is None:
        kind, payload, errors, warnings = parse_excel(content)
    else:
        kind, payload, errors, warnings = "official", official, [], []
    digest = hashlib.sha256(content).hexdigest()
    existing = c.execute(
        "SELECT id FROM drug_imports WHERE sha256=? AND kind=? AND committed_at IS NOT NULL ORDER BY id DESC LIMIT 1",
        (digest, kind),
    ).fetchone()
    if kind == "official":
        exact = c.execute(
            "SELECT id FROM drug_imports WHERE sha256=? AND kind=? AND payload=? AND committed_at IS NOT NULL ORDER BY id DESC LIMIT 1",
            (digest, kind, dump(payload)),
        ).fetchone()
        if exact:
            return exact["id"]
    changes, seen = [], set()
    known_drugs = current(c, "drug")
    identities = {}
    for entity_kind, key, value in payload.get("entities", []):
        tag = (entity_kind, key)
        if tag in seen:
            errors.append(f"Повтор ключа {entity_kind}/{key}; исправьте дубликаты.")
            continue
        seen.add(tag)
        if entity_kind == "adverse":
            raw = value["raw"]
            linked = known_drugs.get(str(raw["ID препарата"]))
            if not linked or normalize(linked["inn"]) != normalize(raw["МНН"]):
                errors.append(
                    f'Реакции, строка {value["row"]}: ID/МНН не совпадает с библиотекой; сначала импортируйте соответствующий препарат.'
                )
        old = c.execute(
            "SELECT payload FROM drug_entities WHERE kind=? AND entity_key=? ORDER BY import_id DESC LIMIT 1",
            tag,
        ).fetchone()
        # INN identity cannot silently change underneath reviewed official facts.
        if (
            entity_kind == "drug"
            and old
            and normalize(json.loads(old[0])["inn"]) != normalize(value["inn"])
        ):
            errors.append(f"{key}: смена МНН требует нового ID препарата.")
        if entity_kind == "drug":
            identity = tuple(normalize(value[k]) for k in ("name", "inn", "form"))
            other = identities.get(identity)
            if other and other != key:
                errors.append(
                    f"{key} и {other}: одинаковые название, МНН и форма под разными ID."
                )
            identities[identity] = key
        changes.append(
            dict(
                kind=entity_kind,
                key=key,
                status=(
                    "new"
                    if not old
                    else ("same" if old[0] == dump(value) else "update")
                ),
                before=json.loads(old[0]) if old else None,
                after=value,
            )
        )
    report = dict(
        errors=errors,
        warnings=warnings,
        changes=changes,
        duplicate_import=existing["id"] if existing else None,
        base=c.execute(
            "SELECT COALESCE(MAX(id),0) FROM drug_imports WHERE committed_at IS NOT NULL"
        ).fetchone()[0],
    )
    iid = c.execute(
        "INSERT INTO drug_imports(kind,filename,sha256,content,payload,report,actor_id,created_at) VALUES(?,?,?,?,?,?,?,?)",
        (
            kind,
            Path(filename).name[:180],
            digest,
            content,
            dump(payload),
            dump(report),
            actor_id,
            now(),
        ),
    ).lastrowid
    return iid


def commit(c, iid):
    item = c.execute("SELECT * FROM drug_imports WHERE id=?", (iid,)).fetchone()
    if item["committed_at"]:
        return
    report = json.loads(item["report"])
    if report["errors"]:
        fail("Импорт содержит ошибки. Исправьте файл и создайте новый предпросмотр.")
    latest = c.execute(
        "SELECT COALESCE(MAX(id),0) FROM drug_imports WHERE committed_at IS NOT NULL"
    ).fetchone()[0]
    if report["base"] != latest:
        raise HTTPException(
            409,
            "Справочник изменился после предпросмотра. Загрузите файл повторно для сравнения.",
        )
    for change in report["changes"]:
        if change["status"] != "same":
            c.execute(
                "INSERT INTO drug_entities VALUES(?,?,?,?)",
                (change["kind"], change["key"], iid, dump(change["after"])),
            )
    c.execute("UPDATE drug_imports SET committed_at=? WHERE id=?", (now(), iid))


def bootstrap(c):
    if c.execute(
        "SELECT 1 FROM drug_bootstrap WHERE version='provided-xlsx-v1'"
    ).fetchone():
        return
    root = Path(__file__).resolve().parents[2] / "docs" / "inputs" / "2026-10-21"
    for filename in [
        "Библиотека препаратов.xlsx",
        "Побочные эффекты по категориям.xlsx",
    ]:
        iid = preview(c, filename, (root / filename).read_bytes(), None)
        commit(c, iid)
    c.execute("INSERT INTO drug_bootstrap VALUES('provided-xlsx-v1')")


def official_payload(c, content, filename, form):
    code = str(form.get("drug_code", ""))
    if code not in current(c, "drug"):
        fail("Выберите препарат из справочника.")
    meta = {
        k: str(form.get(k, "")).strip()
        for k in [
            "source_url",
            "version",
            "registration",
            "product_form",
            "checked_on",
            "valid_until",
        ]
    }
    if not safe_url(meta["source_url"], official=True):
        fail(
            "Укажите HTTPS-ссылку на конкретную запись/документ ГРЛС, а не главную страницу."
        )
    if any(not v or len(v) > 2000 for v in meta.values()):
        fail(
            "Обязательны источник, версия, РУ, конкретная форма и даты проверки/следующей сверки."
        )
    try:
        checked, until = date.fromisoformat(meta["checked_on"]), date.fromisoformat(
            meta["valid_until"]
        )
        if not checked <= date.today() <= until or until > checked + timedelta(
            days=366
        ):
            raise ValueError()
    except ValueError:
        fail(
            "Проверка не может быть в будущем, срок сверки истёк или превышает 366 дней."
        )
    meta["drug_code"] = code
    facts, pairs, entities = [], [], []
    if filename.lower().endswith(".pdf"):
        if not content.startswith(b"%PDF-") or len(content) > MAX_FILE:
            fail("Нужен PDF до 8 МБ.")
        for field in FIELDS:
            value, locator = (
                str(form.get(field, "")).strip(),
                str(form.get(field + "_locator", "")).strip(),
            )
            if value:
                if not locator or len(value) > 15000 or len(locator) > 1000:
                    fail(
                        "У каждого заполненного поля нужна страница/раздел источника; текст до 15000 символов."
                    )
                facts.append(dict(field=field, text=value, locator=locator))
    elif filename.lower().endswith(".xlsx"):
        sheets = spreadsheet(content)
        errors = []
        for _, r in table(
            sheets, "Сведения", "Поле", ["Поле", "Текст", "Раздел/страница"], errors
        ):
            if r["Поле"] not in FIELDS or not r["Текст"] or not r["Раздел/страница"]:
                errors.append("Сведения: неизвестное поле, пустой текст или раздел.")
                continue
            facts.append(
                dict(
                    field=r["Поле"],
                    text=str(r["Текст"]),
                    locator=str(r["Раздел/страница"]),
                )
            )
        for _, r in table(
            sheets,
            "Взаимодействия",
            "Препарат A",
            ["Препарат A", "Препарат B", "Уровень", "Объяснение", "Раздел/страница"],
            errors,
        ):
            a, b = str(r["Препарат A"]), str(r["Препарат B"])
            if (
                a == b
                or code not in (a, b)
                or a not in current(c, "drug")
                or b not in current(c, "drug")
                or r["Уровень"] not in SEVERITIES
                or not r["Объяснение"]
                or not r["Раздел/страница"]
            ):
                errors.append(
                    "Взаимодействия: неверная пара, уровень или отсутствует объяснение/раздел."
                )
                continue
            pairs.append(
                dict(
                    key="|".join(sorted([a, b])),
                    level=r["Уровень"],
                    explanation=str(r["Объяснение"]),
                    locator=str(r["Раздел/страница"]),
                )
            )
        if errors:
            fail("; ".join(errors))
        meta["raw_sheets"] = sheets
    else:
        fail("Поддерживаются PDF инструкции или XLSX по шаблону.")
    if not facts and not pairs:
        fail("Укажите хотя бы одно проверяемое поле или взаимодействие.")
    if len({f["field"] for f in facts}) != len(facts) or len(
        {p["key"] for p in pairs}
    ) != len(pairs):
        fail("Повторяющиеся поля или сочетания в официальной выгрузке.")
    meta.update(facts=facts, pairs=pairs)
    # Official facts have their own event history; importing an Excel never promotes them.
    return dict(meta=meta, entities=entities)


def sources(c, code=None, include_inactive=False):
    result = []
    superseded = set()
    for row in c.execute(
        "SELECT * FROM drug_imports WHERE kind='official' AND committed_at IS NOT NULL ORDER BY id DESC"
    ):
        meta = json.loads(row["payload"])["meta"]
        event = c.execute(
            "SELECT * FROM drug_source_events WHERE import_id=? ORDER BY id DESC LIMIT 1",
            (row["id"],),
        ).fetchone()
        identity = tuple(
            normalize(meta[k]) for k in ("drug_code", "registration", "product_form")
        )
        ever_verified = c.execute(
            "SELECT 1 FROM drug_source_events WHERE import_id=? AND action='verify' LIMIT 1",
            (row["id"],),
        ).fetchone()
        active = bool(
            identity not in superseded
            and event
            and event["action"] == "verify"
            and meta["valid_until"] >= date.today().isoformat()
        )
        # A newer reviewed version retires the old one. Withdrawal/expiry never revives old facts.
        if ever_verified:
            superseded.add(identity)
        if code and meta["drug_code"] != code:
            continue
        if active or include_inactive:
            result.append(
                dict(
                    id=row["id"],
                    filename=row["filename"],
                    sha256=row["sha256"],
                    meta=meta,
                    active=active,
                    event=dict(event) if event else None,
                )
            )
    return result


def search(c, q="", group="", status="", *, include_archived=False):
    from app.services.medicine_reference import enrich
    from app.services import medicine_verification as verification

    drugs = current(c, "drug")
    verified = {s["meta"]["drug_code"] for s in sources(c)}
    reviews = current(c, verification.KIND)
    words = normalize(q).split()
    result = []
    for code, d in drugs.items():
        if d.get("archived") and not include_archived:
            continue
        d = verification.decorate(enrich(d), reviews)
        hay = normalize(
            " ".join(
                str(d.get(k, ""))
                for k in [
                    "name",
                    "inn",
                    "substance",
                    "aliases",
                    "atc",
                    "group",
                    "manufacturer",
                ]
            )
        )
        if not all(w in hay for w in words) or (group and group != d["group"]):
            continue
        d.update(
            verified=code in verified, slug="aviandr" if code == "D01" else code.lower()
        )
        if (
            status == "verified"
            and not d["verified"]
            or status == "ready"
            and not d["can_prescribe"]
            or status == "unverified"
            and d["can_prescribe"]
        ):
            continue
        result.append(d)
    return sorted(result, key=lambda d: normalize(d["name"]))


def check_pairs(c, codes):
    from app.services.medicine_reference import enrich, pair

    codes = list(dict.fromkeys(codes))
    drugs = current(c, "drug")
    if not 2 <= len(codes) <= 20 or any(code not in drugs for code in codes):
        fail("Выберите от 2 до 20 различных препаратов из справочника.")
    draft = current(c, "pair")
    verified = sources(c)
    results = []
    for a, b in combinations(codes, 2):
        key = "|".join(sorted([a, b]))
        known = []
        for s in verified:
            known.extend(
                dict(**p, source=s) for p in s["meta"]["pairs"] if p["key"] == key
            )
        results.append(
            dict(
                a=enrich(drugs[a]),
                b=enrich(drugs[b]),
                known=known,
                candidate=draft.get(key),
                editorial=pair(a, b),
            )
        )
    return results
